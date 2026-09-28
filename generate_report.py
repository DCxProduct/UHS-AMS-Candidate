import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Cover Page
title = doc.add_heading('UHS AMS Candidate System', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_heading('Final Focused UAT Retest', 1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2 = doc.add_heading('Post-Remediation Verification', 2)
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

# Test Environment
doc.add_heading('2. Test Environment', level=1)
env_details = [
    "Date: 27 September 2026",
    "Branch: devdevdev",
    "Commit: 757ecf736010c8ba746184fc746aa61c4ca83a1d",
    "Application URL: http://127.0.0.1:8002/login",
    "Testing Scope: Focused retest of specific defects and security findings post-remediation. Not a full UAT suite."
]
for detail in env_details:
    doc.add_paragraph(detail)

# Executive Summary
doc.add_heading('3. Executive Summary', level=1)
doc.add_paragraph("This focused retest evaluated the latest codebase to verify the resolution of specifically targeted issues from the prior UAT phase. Most issues have been successfully addressed, but one high-severity defect related to payment validation (BUG-003) remains unresolved, throwing a 500 error instead of failing validation gracefully. Therefore, production readiness is not yet achieved.")

# Retest Results
doc.add_heading('4. Retest Results', level=1)
table = doc.add_table(rows=1, cols=7)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'ID'
hdr_cells[1].text = 'Finding'
hdr_cells[2].text = 'Previous Status'
hdr_cells[3].text = 'Current Status'
hdr_cells[4].text = 'Verification Type'
hdr_cells[5].text = 'Evidence'
hdr_cells[6].text = 'Notes'

findings = [
    ("BUG-002", "Search state clear", "FAIL", "PASS", "Browser / Test Verified", "Subagent / Automated Tests", "Clearing search effectively empties query string, resets active search label, and restores default records."),
    ("BUG-003", "Payment amount validation", "FAIL", "FAIL", "Browser / Source Verified", "Browser Logs / Subagent", "UI allows negative values and server validation throws a 500 error (BindingResolutionException) rather than a validation error."),
    ("BUG-004", "Receipt uniqueness", "FAIL", "PASS", "Source / Database Verified", "Source Code", "Database migration explicitly added a unique constraint for receipt_number, accompanied by application-level validation."),
    ("SEC-001", "Dynamic upload validation", "FAIL", "PASS", "Source Verified", "Source Code", "FileUpload config now strictly limits accepted MIME types and maximum file size."),
    ("SEC-002", "Audit coverage", "FAIL", "PASS", "Source Verified", "Source Code", "AuditLogger explicitly logs actions for registrar and cashier roles, expanding beyond admin-only."),
    ("SEC-005", "Configuration", "FAIL", "PASS", "Source Verified", "Source Code", ".env.example utilizes safe placeholders without exposing secrets or debugging features."),
    ("SEC-006", "Storage exposure", "FAIL", "PASS", "Source Verified", "Source Code", "Upload disks for payments and form files have been correctly set to private."),
    ("SEC-007", "Seeded credentials", "FAIL", "PASS", "Source Verified", "Source Code", "AdminUserSeeder forces environment variables for credentials in production context, averting default passwords."),
    ("SEC-008", "Automated tests", "BLOCKED", "PASS", "Automated Test Verified", "Test Output", "PHP dependencies were successfully resolved, and targeted tests passed."),
    ("BUG-001", "Candidate search exact/no-result", "PASS", "PASS", "Browser / Test Verified", "Subagent / Tests", "Regression passed. Search continues to perform as expected."),
    ("BUG-005", "Audit Logs 500", "PASS", "PASS", "Automated Test Verified", "Automated Tests", "Regression passed. Audit logs load securely without unhandled exceptions."),
    ("SEC-004", "Registrar permissions", "PASS", "PASS", "Browser Verified", "Browser Subagent", "Regression passed. Unauthorized access explicitly yields 403 Forbidden.")
]

for item in findings:
    row_cells = table.add_row().cells
    for i, text in enumerate(item):
        row_cells[i].text = text

# Remaining Issues
doc.add_heading('5. Remaining Issues', level=1)
doc.add_paragraph("BUG-003: Payment amount validation. The validation logic for the 'amount_kh' field uses an improper closure that causes an unhandled 500 Internal Server Error (BindingResolutionException) instead of a standard validation rejection message. In addition, removing the client-side JavaScript regex allows negative inputs to be typed into the UI.")

# Screenshot Evidence
doc.add_heading('6. Screenshot Evidence', level=1)
doc.add_paragraph("Relevant screenshot evidence has been captured and archived by the browser subagent, demonstrating successful search clearance and the 500 error encountered during payment submission.")

# Production Readiness
doc.add_heading('7. Production Readiness', level=1)
doc.add_paragraph("Production Decision: RETEST REQUIRED")
doc.add_paragraph("While significant progress has been made, the remaining high-severity technical finding (BUG-003) regarding payment application prevents a safe production deployment.")

# Sign-Off
doc.add_heading('8. Sign-Off', level=1)
doc.add_paragraph("Focused Remediation Retest: RETEST REQUIRED")
doc.add_paragraph("Prepared By: Antigravity UAT Engineer")

doc.save('UHS_AMS_Final_Focused_UAT.docx')
